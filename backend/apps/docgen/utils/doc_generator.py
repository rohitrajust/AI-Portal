import re
import tempfile

DOCUMENT_TEMPLATES = {
    "BRD": [
        "Executive Summary",
        "Business Objectives",
        "Scope (In-Scope / Out-of-Scope)",
        "Stakeholders",
        "Functional Requirements",
        "Non-Functional Requirements",
        "Assumptions & Constraints",
        "Risks",
        "Approval",
    ],
    "Mail Body": [
        "Subject",
        "Greeting",
        "Purpose / Context",
        "Key Points",
        "Action Items / Next Steps",
        "Closing",
    ],
    "User Story": [
        "Title",
        "User Story (As a / I want / So that)",
        "Acceptance Criteria",
        "Definition of Done",
        "Additional Notes",
    ],
    "Proposal Document": [
        "Title & Overview",
        "Problem Statement",
        "Proposed Solution",
        "Scope of Work",
        "Timeline",
        "Team / Resources",
        "Benefits",
        "Conclusion / Call to Action",
    ],
}


def build_docx(doc_type: str, document_text: str) -> str:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(doc_type, level=0)

    for line in document_text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        heading_match = re.match(r"^#{1,3}\s+(.*)", stripped)
        bullet_match = re.match(r"^[-*]\s+(.*)", stripped)
        if heading_match:
            doc.add_heading(heading_match.group(1), level=1)
        elif bullet_match:
            doc.add_paragraph(bullet_match.group(1), style="List Bullet")
        else:
            p = doc.add_paragraph(stripped)
            p.style.font.size = Pt(11)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(tmp.name)
    return tmp.name
